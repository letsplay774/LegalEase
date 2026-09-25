import io
import re
from docx import Document
from docx.shared import Pt, RGBColor
from fpdf import FPDF

def sanitize_text(text: str) -> str:
    """Removes markdown bold/italic formatting characters."""
    clean = re.sub(r'[*_]', '', text)
    return clean.strip()

def format_html_preview(text: str) -> str:
    """Formats generator text into basic HTML for Streamlit preview."""
    lines = text.split("\n")
    html_lines = []
    for line in lines:
        line_str = line.strip()
        if not line_str:
            continue
        if line_str.startswith("#"):
            clean_heading = line_str.lstrip("#").strip()
            html_lines.append(f"<h3 style='color: #4CAF50; margin-top: 15px;'>{clean_heading}</h3>")
        else:
            html_lines.append(f"<p style='margin-bottom: 8px; line-height: 1.5;'>{line_str}</p>")
    return "".join(html_lines)

def format_docx(text: str, document_type: str = "Legal Document") -> bytes:
    """Generates a downloadable .docx binary buffer."""
    doc = Document()
    
    # doc.add_heading returns a Paragraph object directly
    heading = doc.add_heading(document_type, level=1)
    
    # Access .runs directly from the Paragraph object
    if heading.runs:
        heading.runs[0].font.size = Pt(20)
        heading.runs[0].font.color.rgb = RGBColor(30, 41, 59)
    
    doc.add_paragraph("")  # Vertical spacing
    
    for line in text.split("\n"):
        line_str = sanitize_text(line)
        if line_str:
            doc.add_paragraph(line_str)
            
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def format_pdf(text: str, document_type: str = "Legal Document") -> bytes:
    """Generates a downloadable .pdf binary buffer."""
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, document_type, ln=True, align="C")
    pdf.ln(10)
    
    pdf.set_font("Helvetica", size=11)
    for line in text.split("\n"):
        line_str = sanitize_text(line)
        if line_str:
            clean_line = line_str.encode('latin-1', 'replace').decode('latin-1')
            pdf.multi_cell(0, 7, clean_line)
            pdf.ln(2)
            
    return bytes(pdf.output())